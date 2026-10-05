import os
from io import BytesIO
from datetime import datetime
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
from reportlab.platypus import KeepTogether
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ─── PDF Generation ────────────────────────────────────────────────────────────

BRAND_PURPLE = colors.HexColor('#6366f1')
BRAND_DARK   = colors.HexColor('#0f172a')
BRAND_SLATE  = colors.HexColor('#1e293b')
BRAND_LIGHT  = colors.HexColor('#e2e8f0')
STATUS_COLORS = {
    'open':        colors.HexColor('#6366f1'),
    'in_progress': colors.HexColor('#f59e0b'),
    'resolved':    colors.HexColor('#10b981'),
    'closed':      colors.HexColor('#64748b'),
}
PRIORITY_COLORS = {
    'low':      colors.HexColor('#22c55e'),
    'medium':   colors.HexColor('#f59e0b'),
    'high':     colors.HexColor('#ef4444'),
    'critical': colors.HexColor('#dc2626'),
}


def generate_tickets_pdf(tickets, filters=None, stats=None):
    """Generate a PDF report for a list of tickets. Returns BytesIO buffer."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5*inch, bottomMargin=0.5*inch,
                            leftMargin=0.6*inch, rightMargin=0.6*inch)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle('Title', fontSize=22, textColor=BRAND_PURPLE,
                                 spaceAfter=4, fontName='Helvetica-Bold', alignment=TA_LEFT)
    sub_style = ParagraphStyle('Sub', fontSize=10, textColor=colors.HexColor('#94a3b8'),
                               spaceAfter=2, fontName='Helvetica')
    story.append(Paragraph("IT Help Desk — Ticket Report", title_style))
    story.append(Paragraph(f"Generated: {datetime.utcnow().strftime('%B %d, %Y at %H:%M UTC')} | Total Tickets: {len(tickets)}", sub_style))
    story.append(HRFlowable(width="100%", thickness=1, color=BRAND_PURPLE, spaceAfter=12))

    # Stats summary table
    if stats:
        stat_data = [
            ['Open', 'In Progress', 'Resolved', 'Closed'],
            [str(stats.get('open', 0)), str(stats.get('in_progress', 0)),
             str(stats.get('resolved', 0)), str(stats.get('closed', 0))]
        ]
        stat_table = Table(stat_data, colWidths=[1.6*inch]*4)
        stat_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), BRAND_SLATE),
            ('TEXTCOLOR', (0,0), (-1,0), BRAND_LIGHT),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 9),
            ('BACKGROUND', (0,1), (-1,1), colors.HexColor('#f1f5f9')),
            ('FONTSIZE', (0,1), (-1,1), 16),
            ('FONTNAME', (0,1), (-1,1), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0,1), (-1,1), [colors.HexColor('#f8fafc')]),
            ('BOX', (0,0), (-1,-1), 1, BRAND_PURPLE),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ]))
        story.append(stat_table)
        story.append(Spacer(1, 16))

    # Tickets table
    header = ['Ticket #', 'Title', 'Category', 'Priority', 'Status', 'Assigned To', 'Created']
    rows = [header]
    for t in tickets:
        rows.append([
            t.ticket_number,
            (t.title[:40] + '...') if len(t.title) > 40 else t.title,
            t.category,
            t.priority.upper(),
            t.status.replace('_', ' ').title(),
            t.assignee.name if t.assignee else '—',
            t.created_at.strftime('%Y-%m-%d'),
        ])

    col_widths = [1.0*inch, 2.1*inch, 0.85*inch, 0.75*inch, 0.85*inch, 1.05*inch, 0.85*inch]
    ticket_table = Table(rows, colWidths=col_widths, repeatRows=1)
    ticket_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), BRAND_DARK),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')]),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(ticket_table)

    doc.build(story)
    buffer.seek(0)
    return buffer


# ─── Excel Generation ──────────────────────────────────────────────────────────

def generate_tickets_excel(tickets):
    """Generate Excel report. Returns BytesIO buffer."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tickets Report"

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_align = Alignment(horizontal='center', vertical='center')
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    status_fills = {
        'open':        PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid"),
        'in_progress': PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid"),
        'resolved':    PatternFill(start_color="10B981", end_color="10B981", fill_type="solid"),
        'closed':      PatternFill(start_color="64748B", end_color="64748B", fill_type="solid"),
    }
    priority_fills = {
        'low':      PatternFill(start_color="22C55E", end_color="22C55E", fill_type="solid"),
        'medium':   PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid"),
        'high':     PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid"),
        'critical': PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid"),
    }

    # Title row
    ws.merge_cells('A1:J1')
    ws['A1'] = f'IT Help Desk — Tickets Report | Generated: {datetime.utcnow().strftime("%B %d, %Y")}'
    ws['A1'].font = Font(bold=True, size=14, color="6366F1")
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    # Headers
    headers = ['Ticket #', 'Title', 'Category', 'Priority', 'Status', 'Assigned To', 'Created By', 'Created At', 'Due Date', 'Resolution Time (hrs)']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[2].height = 24

    # Data rows
    for row_idx, ticket in enumerate(tickets, 3):
        data = [
            ticket.ticket_number,
            ticket.title,
            ticket.category,
            ticket.priority.upper(),
            ticket.status.replace('_', ' ').title(),
            ticket.assignee.name if ticket.assignee else 'Unassigned',
            ticket.creator.name if ticket.creator else 'Unknown',
            ticket.created_at.strftime('%Y-%m-%d %H:%M'),
            ticket.due_date.strftime('%Y-%m-%d') if ticket.due_date else '—',
            ticket.resolution_time_hours() or '—',
        ]
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical='center')
            # Color status cell
            if col == 5:
                fill = status_fills.get(ticket.status)
                if fill:
                    cell.fill = fill
                    cell.font = Font(color='FFFFFF', bold=True)
            # Color priority cell
            if col == 4:
                fill = priority_fills.get(ticket.priority)
                if fill:
                    cell.fill = fill
                    cell.font = Font(color='FFFFFF', bold=True)
            # Alternating row color
            if col not in (4, 5) and row_idx % 2 == 0:
                cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Column widths
    col_widths = [14, 40, 16, 12, 16, 20, 20, 20, 14, 22]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = 'A3'
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
